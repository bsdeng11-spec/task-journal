"""
Excel Exporter Module.
Generates the exact Master Task Journal Spreadsheet matching the company format from Image 1.
Columns: Name | Task | Description | Date | General hours | Executor | MP#
Features dark navy banner header (Rows 1 & 2), dark navy table headers with AutoFilter (Row 3),
clean data rows grouped by Executor with blank row spacing, and a clean Summary Dashboard tab.
"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY_FILL_COLOR = "0A1931"  # Deep dark navy matching company template
WHITE_COLOR = "FFFFFF"

def format_date_mdy(date_val):
    if not date_val:
        return ""
    if isinstance(date_val, datetime):
        return f"{date_val.month}/{date_val.day}/{date_val.year}"
    try:
        dt = datetime.strptime(str(date_val).strip(), "%Y-%m-%d")
        return f"{dt.month}/{dt.day}/{dt.year}"
    except Exception:
        return str(date_val)

def generate_master_excel(period_title, year_label, all_members, tasks_by_member):
    """
    Builds the exact master report workbook matching company format from Image 1.
    all_members: list of dicts [{'id': 1, 'name': 'Robel', 'mp_number': '9826'}, ...]
    tasks_by_member: dict {member_id: [task_dict, ...]}
    """
    wb = Workbook()

    # -------------------------------------------------------------
    # TAB 1: Master Task Journal (Exact layout from Image 1)
    # -------------------------------------------------------------
    ws = wb.active
    ws.title = "Master Task Journal"

    # Styling definitions
    navy_fill = PatternFill(start_color=NAVY_FILL_COLOR, end_color=NAVY_FILL_COLOR, fill_type="solid")
    font_banner = Font(name="Arial", size=11, bold=True, color=WHITE_COLOR)
    font_header = Font(name="Arial", size=11, bold=True, color=WHITE_COLOR)
    font_data = Font(name="Arial", size=10)

    thin_border_side = Side(style='thin', color='E5E7EB')
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Row 1: Report Title Banner (Cols A-G filled with dark navy)
    ws.append([f"Report_{year_label}", "", "", "", "", "", ""])
    ws.row_dimensions[1].height = 26
    for c in range(1, 8):
        cell = ws.cell(row=1, column=c)
        cell.fill = navy_fill
        cell.font = font_banner
        cell.alignment = Alignment(vertical="center", horizontal="left" if c == 1 else "center")

    # Row 2: Period Label Banner (e.g. SEP 16-30)
    ws.append([str(period_title).upper(), "", "", "", "", "", ""])
    ws.row_dimensions[2].height = 22
    for c in range(1, 8):
        cell = ws.cell(row=2, column=c)
        cell.fill = navy_fill
        cell.font = font_banner
        cell.alignment = Alignment(vertical="center", horizontal="left" if c == 1 else "center")

    # Row 3: Table Headers matching Image 1
    headers = ["Name", "Task", "Description", "Date", "General hours", "Executor", "MP#"]
    ws.append(headers)
    ws.row_dimensions[3].height = 24
    for c in range(1, 8):
        cell = ws.cell(row=3, column=c)
        cell.fill = navy_fill
        cell.font = font_header
        cell.alignment = Alignment(
            vertical="center",
            horizontal="center" if c in [1, 4, 5, 7] else "left"
        )

    grand_total_hours = 0.0

    # Data rows grouped by Executor
    for member in all_members:
        m_id = member['id']
        m_tasks = tasks_by_member.get(m_id, [])
        if not m_tasks:
            continue

        for t in m_tasks:
            proj = t.get('project_code', '') or ''
            task_type = t.get('task_type', '') or ''
            desc = t.get('description', '') or ''
            date_raw = t.get('date', '') or ''
            date_display = format_date_mdy(date_raw)
            hrs_val = t.get('hours')
            is_special = t.get('is_special', 'Work')

            # Weekend or Leave check (Vacation, Sick, Emergency, Other, Holiday)
            desc_clean = desc.strip().lower()
            is_leave_or_weekend = (
                is_special in ['Weekend', 'Holiday', 'Vacation', 'Sick', 'Emergency', 'Other'] or
                any(k in desc_clean for k in ['weekend', 'holiday', 'vacation', 'vaccation', 'sick', 'emergency', 'other leave']) or
                str(hrs_val).strip() == '-'
            )

            if is_leave_or_weekend:
                leave_label = desc if desc else (
                    "Weekend" if is_special == 'Weekend' else (
                        "Vacation" if is_special == 'Vacation' else (
                            "Sick Leave" if is_special == 'Sick' else (
                                "Emergency Leave" if is_special == 'Emergency' else (
                                    "Other Leave" if is_special == 'Other' else "Leave"
                                )
                            )
                        )
                    )
                )
                row_data = [
                    "",
                    "",
                    leave_label,
                    date_display,
                    "-",
                    member['name'],
                    member['mp_number']
                ]
            else:
                if hrs_val not in [None, "", "-"]:
                    try:
                        numeric_hrs = float(hrs_val)
                        grand_total_hours += numeric_hrs
                        hrs_display = int(numeric_hrs) if numeric_hrs.is_integer() else round(numeric_hrs, 2)
                    except (ValueError, TypeError):
                        hrs_display = hrs_val
                else:
                    hrs_display = ""

                row_data = [
                    proj,
                    task_type,
                    desc,
                    date_display,
                    hrs_display,
                    member['name'],
                    member['mp_number']
                ]

            ws.append(row_data)
            curr_row = ws.max_row
            ws.row_dimensions[curr_row].height = 20

            for col_idx in range(1, 8):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = font_data
                cell.border = border_cell
                if col_idx in [1, 4, 7]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 5:
                    cell.alignment = Alignment(horizontal="right" if row_data[4] != "-" else "center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # 3 blank rows between executors as shown in Image 1
        for _ in range(3):
            ws.append(["", "", "", "", "", "", ""])
            curr_blank = ws.max_row
            ws.row_dimensions[curr_blank].height = 18

    # Apply AutoFilter across columns A to G on header row 3
    last_row = max(ws.max_row, 3)
    ws.auto_filter.ref = f"A3:G{last_row}"

    # Column widths matching screenshot proportions
    col_widths = {
        'A': 14,  # Name
        'B': 14,  # Task
        'C': 48,  # Description
        'D': 13,  # Date
        'E': 14,  # General hours
        'F': 15,  # Executor
        'G': 10   # MP#
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # -------------------------------------------------------------
    # TAB 2: Summary Dashboard
    # -------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Summary Dashboard")
    ws_summary.append(["TASK JOURNAL 15-DAY SUMMARY DASHBOARD"])
    ws_summary.cell(row=1, column=1).font = Font(name="Arial", size=14, bold=True, color="0A1931")
    ws_summary.append([f"Period: {period_title} ({year_label})", f"Grand Total Hours: {round(grand_total_hours, 2)} hrs"])
    ws_summary.append([])

    # Table of Team Members
    sum_headers = ["#", "Executor", "MP Number", "Logged Hours", "Target Hours (15-Day)", "Status"]
    ws_summary.append(sum_headers)
    hdr_row = ws_summary.max_row
    fill_table_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_total = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=10)

    for c_idx in range(1, len(sum_headers) + 1):
        cell = ws_summary.cell(row=hdr_row, column=c_idx)
        cell.font = font_bold
        cell.fill = fill_table_header
        cell.border = border_cell
        cell.alignment = Alignment(horizontal="center", vertical="center")

    idx = 1
    for member in all_members:
        m_id = member['id']
        m_tasks = tasks_by_member.get(m_id, [])
        m_hours = 0.0
        for t in m_tasks:
            h = t.get('hours')
            if h not in [None, "", "-"]:
                try:
                    m_hours += float(h)
                except (ValueError, TypeError):
                    pass

        status = "Complete" if m_hours >= 99.0 else ("In Progress" if m_hours > 0 else "Not Started")

        status_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") if status == "Complete" else (
            PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") if status == "In Progress" else
            PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        )

        ws_summary.append([idx, member['name'], member['mp_number'], round(m_hours, 2), 99.0, status])
        curr = ws_summary.max_row
        for col_idx in range(1, 7):
            c = ws_summary.cell(row=curr, column=col_idx)
            c.border = border_cell
            c.font = font_regular
            if col_idx in [1, 3, 4, 5]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.fill = status_fill
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
        idx += 1

    # Total row on summary
    ws_summary.append(["", "TOTAL", "", round(grand_total_hours, 2), len(all_members) * 99.0, ""])
    curr = ws_summary.max_row
    for col_idx in range(1, 7):
        c = ws_summary.cell(row=curr, column=col_idx)
        c.font = font_bold
        c.fill = fill_total
        c.border = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))
        if col_idx in [2, 4]:
            c.alignment = Alignment(horizontal="center" if col_idx == 2 else "right")

    for col in ws_summary.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
